"""Text extraction for contact-bearing document resources."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import pandas as pd
from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".csv", ".xls", ".xlsx"}

def is_document_url(url: str, content_type: str = "") -> bool:
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    return suffix in SUPPORTED_EXTENSIONS or content_type.lower() in {
        "application/pdf", "application/msword",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "text/csv", "application/csv", "application/vnd.ms-excel",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }

def extract_document_text(body: bytes, url: str, content_type: str = "") -> str:
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    if suffix == ".pdf" or content_type == "application/pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(body)).pages)
    if suffix in {".docx"}:
        doc = Document(BytesIO(body))
        parts = [p.text for p in doc.paragraphs]
        parts += [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
        return "\n".join(p for p in parts if p.strip())
    if suffix in {".csv"} or content_type in {"text/csv", "application/csv"}:
        return pd.read_csv(BytesIO(body), dtype=str, keep_default_na=False).to_csv(index=False)
    if suffix in {".xls", ".xlsx"} or "spreadsheetml" in content_type or content_type == "application/vnd.ms-excel":
        sheets = pd.read_excel(BytesIO(body), sheet_name=None, dtype=str)
        return "\n".join(f"[{name}]\n{frame.to_csv(index=False)}" for name, frame in sheets.items())
    if suffix == ".doc":
        raise ValueError("legacy .doc extraction is not supported; convert to .docx or PDF")
    return ""
